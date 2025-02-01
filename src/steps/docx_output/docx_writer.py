import shutil
from docx import Document
from docx.shared import Inches
import requests
import packaging.version
from io import BytesIO
from docx.shared import RGBColor
import os
import matplotlib.pyplot as plt


def add_heading(doc, text, level=1):
    doc.add_heading(text, level)


def add_sub_heading(doc, text, level=2):
    doc.add_heading(text, level)


from docx.shared import Pt  # Make sure to import Pt


def add_paragraph(doc, text, style=None, bold=False, small_font=False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)

    if bold:
        run.bold = True
    if small_font:
        run.font.size = Pt(8)  # Use Pt() for point measurements

    return paragraph


def add_table(doc, data, style="Table Grid"):
    rows, cols = len(data), len(data[0])
    table = doc.add_table(rows=rows, cols=cols)
    table.style = style
    for i, row in enumerate(data):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = cell_text
            if i == 0:  # Make the heading row bold and blue
                run = cell.paragraphs[0].runs[0]
                run.bold = True
                run.font.color.rgb = RGBColor(0, 0, 255)  # Blue color


def add_image(doc, image_path, width=Inches(4.0)):
    doc.add_picture(image_path, width=width)


def add_chart(doc, chart_func, *args, **kwargs):
    fig, ax = plt.subplots()
    chart_func(ax, *args, **kwargs)
    chart_image = BytesIO()
    plt.savefig(chart_image, format="png")
    chart_image.seek(0)
    doc.add_picture(chart_image, width=Inches(4.0))
    plt.close(fig)


def create_reference_doc():
    ref_doc = Document()
    add_heading(ref_doc, "References", level=1)
    add_paragraph(
        ref_doc,
        "1. Author A. (Year). Title of the paper. Journal Name, Volume(Issue), pages.",
    )
    add_paragraph(ref_doc, "2. Author B. (Year). Title of the book. Publisher.")
    return ref_doc


def save_document(doc, path):
    doc.save(path)


def append_docs(main_doc, ref_doc_path):
    with open(ref_doc_path, "rb") as ref_file:
        ref_doc = Document(BytesIO(ref_file.read()))
    for element in ref_doc.element.body:
        main_doc.element.body.append(element)


thesis_path = "/Users/laxmanbhattarai/projects/personal/mba/thesis_v2/src/steps/docx_output/thesis.docx"

ref_path = "/Users/laxmanbhattarai/projects/personal/mba/thesis_v2/src/steps/docx_output/refs.docx"


def open_thesis():
    desktop_path ="/Users/laxmanbhattarai/Desktop/thesis.docx"
    shutil.copy(thesis_path, desktop_path)
    os.system(f"open {thesis_path}")


def save_document(doc, path):
    doc.save(path)


def clear_thesis():
    if os.path.exists(thesis_path):
        os.remove(thesis_path)
    if os.path.exists(ref_path):
        os.remove(ref_path)


def thesis_body():
    path = thesis_path
    if not os.path.exists(path):
        return Document()
    with open(path, "rb") as file:
        return Document(BytesIO(file.read()))


def thesis_ref():
    path = ref_path
    if not os.path.exists(path):
        return Document()
    with open(path, "rb") as file:
        return Document(BytesIO(file.read()))


def write_dummy_docx():
    body = thesis_body()
    add_heading(body, "Document Title", 0)
    add_paragraph(body, "This is a simple paragraph.")
    add_paragraph(body, "First item in the list", style="List Bullet")
    add_paragraph(body, "Second item in the list", style="List Bullet")
    add_paragraph(body, "First item in the numbered list", style="List Number")
    add_paragraph(body, "Second item in the numbered list", style="List Number")

    table_data = [["Name", "Age"], ["John Doe", "30"]]
    add_table(body, table_data)

    add_chart(
        body,
        lambda ax: ax.pie(
            [15, 30, 45, 10],
            labels=["A", "B", "C", "D"],
            autopct="%1.1f%%",
            startangle=90,
        ),
    )
    add_chart(body, lambda ax: ax.scatter([1, 2, 3, 4, 5], [10, 20, 25, 30, 40]))
    add_chart(body, lambda ax: ax.boxplot([20, 30, 40, 50, 60, 70, 80, 90, 100]))
    # plot([20, 30, 40, 50, 60, 70, 80, 90, 100], positions=[1], widths=0.6)
    # ax.boxplot([15, 25, 35, 45, 55, 65, 75, 85, 95], positions=[2], widths=0.6)
    # ax.set_xticks([1, 2])
    # ax.set_xticklabels(['Boxplot 1', 'Boxplot 2'])
    add_chart(body, lambda ax: ax.plot([1, 2, 3, 4, 5], [10, 20, 30, 40, 50]))
    fig, axs = plt.subplots(1, 3, figsize=(12, 4))
    data1 = [20, 30, 40, 50, 60, 70, 80, 90, 100]
    data2 = [15, 25, 35, 45, 55, 65, 75, 85, 95]
    data3 = [10, 20, 30, 40, 50, 60, 70, 80, 90]

    axs[0].boxplot(data1)
    axs[0].set_title('Boxplot 1')

    axs[1].boxplot(data2)
    axs[1].set_title('Boxplot 2')

    axs[2].boxplot(data3)
    axs[2].set_title('Boxplot 3')

    plt.tight_layout()
    chart_image = BytesIO()
    plt.savefig(chart_image, format="png")
    chart_image.seek(0)
    body.add_picture(chart_image, width=Inches(6.0))
    plt.close(fig)
    reference = thesis_ref()
    long_table_data = [["Index", "Description"]] + [
        [str(i), f"Description for item {i}"] for i in range(1, 21)
    ]
    add_table(reference, long_table_data)

    save_document(reference, ref_path)

    save_document(body, thesis_path)

    append_docs(body, ref_path)
    save_document(body, thesis_path)
